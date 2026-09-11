# -*- coding: utf-8 -*-
"""ProjeTakip.xlsm -- personelin oneri takip kitabinin sayfa kurulumu.

Tek ekran: Takip. Personel oneri numarasini ve sicil numarasini yazar,
"Sorgula"ya basar; guncel durum ve durum gecmisi ekrana gelir.

Bu kitapta VERI SAYFASI YOKTUR ve olmamalidir. Kitap parolasizdir; depo
buraya kopyalansaydi herkesin onerisi herkesin elinde olurdu. Sorgu depoyu
gizli bir Excel orneginde okur ve yalnizca istenen onerinin satirlarini
getirir (bkz. modTakip.bas).
"""

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName

import uret_ortak as u
from tasarim import DUGME, RENK

SAYFA_TAKIP = "Takip"

# Oneri formuyla ayni izgara -- A: kenar, B: etiket, C-F: alan, G: kenar
SUTUNLAR = [2.2, 24.0, 20.0, 20.0, 20.0, 14.0, 2.2]
SON_SUTUN = len(SUTUNLAR)
C_ETIKET = 2
C_ALAN_BAS = 3
C_ALAN_SON = 6

# --- Durum gecmisi tablosu (modTakip.bas'taki sabitlerle ayni olmali) ------
GECMIS_BASLIK = 30
GECMIS_ILK = 31                  # modTakip.GECMIS_ILK_SATIR
GECMIS_ADET = 10                 # modTakip.GECMIS_AZAMI

# Yan yana duran dugmeler AYNI hucreye capalanir (bkz. uret_oneri).
DUGME_YERLERI = {"sorgula": "C16", "temizle": "C16"}
DUGME_KAYDIR = {"temizle": DUGME["GENISLIK"] + DUGME["ARALIK"]}


def _ad(ad, sutun, satir):
    return DefinedName(ad, attr_text=f"'{SAYFA_TAKIP}'!${get_column_letter(sutun)}${satir}")


def _giris_alani(ws, wb, satir, etiket_metni, ad):
    u.etiket(ws, satir, C_ETIKET, etiket_metni, zorunlu=True)
    u.form_alani(ws, wb, satir, C_ALAN_BAS, C_ALAN_SON, ad, SAYFA_TAKIP)
    # Metin bicimi: "010045" gibi bir sicil basindaki sifiri kaybetmesin.
    for c in range(C_ALAN_BAS, C_ALAN_SON + 1):
        ws.cell(row=satir, column=c).number_format = "@"


def _takip_sayfasi(wb):
    ws = wb.create_sheet(SAYFA_TAKIP)
    u.sayfa_hazirla(ws, SUTUNLAR, son_satir=45)
    u.masthead(ws, SON_SUTUN, "Öneri Takibi")

    u.bosluk(ws, 4, 12.0)

    # Durum bandi: sorgunun sonucunu (bulundu / bulunamadi) tek satirda soyler.
    ws.row_dimensions[5].height = 26.0
    bant = u.birlestir(ws, 5, C_ETIKET, 5, C_ALAN_SON)
    bant.alignment = u.hiza("left", "center", girinti=1)
    wb.defined_names.add(_ad("tk_bant", C_ETIKET, 5))

    u.bosluk(ws, 6, 12.0)

    # --- 1. Sorgu ---------------------------------------------------------
    u.bolum_basligi(ws, 7, C_ETIKET, C_ALAN_SON, "1  ·  SORGU",
                    "Öneri numaranızı ve öneriyi gönderirken yazdığınız "
                    "sicil numarasını girin.")
    u.bosluk(ws, 9, 10.0)

    _giris_alani(ws, wb, 10, "Öneri No", "tk_no")
    u.ipucu_satiri(ws, 11, C_ALAN_BAS, C_ALAN_SON,
                   "Gönderimde verilen numara. Örnek: “PRJ-2026-0001”")
    u.bosluk(ws, 12, 6.0)
    _giris_alani(ws, wb, 13, "Sicil No", "tk_sicil")
    u.ipucu_satiri(ws, 14, C_ALAN_BAS, C_ALAN_SON,
                   "Numara ile sicil birlikte tutmazsa sonuç gösterilmez.")

    u.bosluk(ws, 15, 12.0)
    u.bosluk(ws, 16, 44.0)      # Sorgula / Temizle dugmeleri
    u.bosluk(ws, 17, 14.0)
    u.ayrac(ws, 18, C_ETIKET, C_ALAN_SON)
    u.bosluk(ws, 19, 16.0)

    # --- 2. Sonuc (salt okunur) --------------------------------------------
    u.bolum_basligi(ws, 20, C_ETIKET, C_ALAN_SON, "2  ·  ÖNERİNİZİN DURUMU",
                    "Bilgiler sorgu anında okunur; ekranı açık bıraktıysanız "
                    "yeniden sorgulayın.")
    u.bosluk(ws, 22, 8.0)

    u.okuma_alani(ws, wb, SAYFA_TAKIP, 23, 2, "Öneri No", "tk_sonuc_no", 3, 3)
    u.okuma_alani(ws, wb, SAYFA_TAKIP, 23, 5, "Gönderim Tarihi", "tk_tarih", 6, 6)
    u.okuma_alani(ws, wb, SAYFA_TAKIP, 24, 2, "Öneri Başlığı", "tk_baslik", 3, 6)
    durum = u.okuma_alani(ws, wb, SAYFA_TAKIP, 25, 2, "Güncel Durum", "tk_durum",
                          3, 3, yukseklik=26.0)
    durum.alignment = u.hiza("center", "center")
    u.okuma_alani(ws, wb, SAYFA_TAKIP, 25, 5, "Son Güncelleme", "tk_guncelleme",
                  6, 6, yukseklik=26.0)

    # Sonuc hucreleri METIN bicimli: "04.09.2026" aksi halde tarihe cevrilir.
    for hucre in ("C23", "F23", "C24", "C25", "F25"):
        ws[hucre].number_format = "@"
    u.durum_kosullu_bicim(ws, "C25")

    u.bosluk(ws, 26, 20.0)

    # --- 3. Durum gecmisi ----------------------------------------------------
    u.bolum_basligi(ws, 27, C_ETIKET, C_ALAN_SON, "3  ·  DURUM GEÇMİŞİ",
                    "En yeni değişiklik üstte.")
    u.bosluk(ws, 29, 6.0)

    son = GECMIS_ILK + GECMIS_ADET - 1
    u.tablo_basligi(ws, GECMIS_BASLIK, 2, 6, ["Tarih", "Durum", "Açıklama"],
                    ortali=(3,))
    u.tablo_govde_stili(ws, GECMIS_ILK, son, 2, 6, ortali=(3,))

    # Aciklama sutunu uc sutun genisliginde: durumun anlami tek satirda okunsun.
    for r in range(GECMIS_BASLIK, son + 1):
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
    for r in range(GECMIS_ILK, son + 1):
        for c in (2, 3, 4):
            ws.cell(row=r, column=c).number_format = "@"
    u.durum_kosullu_bicim(ws, f"C{GECMIS_ILK}:C{son}")

    u.bosluk(ws, son + 1, 14.0)
    u.not_satiri(ws, son + 2, C_ETIKET, C_ALAN_SON,
                 "Karar ayrıntıları için öneri numaranızla Değerlendirme "
                 "ekibine başvurabilirsiniz.")
    u.not_satiri(ws, son + 3, C_ETIKET, C_ALAN_SON,
                 "Dosya açıldığında üstte bir güvenlik uyarısı çıkarsa "
                 "“İçeriği Etkinleştir” düğmesine basın.")

    ws.sheet_state = "visible"
    return ws


def kitap_uret(hedef_yol):
    wb = Workbook()
    wb.remove(wb.active)

    ws = _takip_sayfasi(wb)
    ws.sheet_properties.tabColor = RENK["DERIN_LACIVERT"]

    wb.active = wb.index(ws)
    wb.save(hedef_yol)
    return hedef_yol
