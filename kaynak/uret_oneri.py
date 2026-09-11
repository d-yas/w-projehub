# -*- coding: utf-8 -*-
"""ProjeOneri.xlsm -- personel gonderim kitabinin sayfa kurulumu.

Bu betik yalnizca gorsel/yapisal iskeleti kurar (.xlsx taslak). VBA modulleri,
dugme sekilleri ve korumalar com_kurulum.py asamasinda eklenir.
"""

from openpyxl import Workbook
from openpyxl.worksheet.datavalidation import DataValidation

import uret_ortak as u
from tasarim import DUGME, RENK

SAYFA_GIRIS = "Giriş"
SAYFA_FORM = "Öneri Formu"

# A: kenar boslugu, B: etiket, C-F: giris alani, G: kenar boslugu
SUTUNLAR = [2.2, 24.0, 20.0, 20.0, 20.0, 14.0, 2.2]
SON_SUTUN = len(SUTUNLAR)
C_ETIKET = 2
C_ALAN_BAS = 3
C_ALAN_SON = 6

# Dugmelerin uzerine oturacagi hucreler (COM asamasi bu adresleri kullanir).
# Yan yana duran dugmeler AYNI hucreye capalanir; aralarindaki mesafe sutun
# genisligine degil DUGME["ARALIK"]a baglidir -- boylece serit, sutunlar
# degistiginde dagilmaz.
DUGME_YERLERI = {
    "giris": "C10",
    "gonder": "C32",
    "temizle": "C32",
    "cikis": "F5",
}

DUGME_KAYDIR = {
    "temizle": DUGME["GENISLIK"] + DUGME["ARALIK"],
}


def _giris_sayfasi(wb):
    """Giris ekrani: zeminin ustunde tek bir beyaz kart.

    Kartin icinde yalnizca uc sey var -- baslik, tek cumlelik aciklama ve
    dugme. Web'deki giris sayfalarinin deseni: bir sey yapilacaksa nerede
    yapilacagi tartisilmaz.
    """
    ws = wb.create_sheet(SAYFA_GIRIS)
    u.sayfa_hazirla(ws, SUTUNLAR, son_satir=18)
    u.masthead(ws, SON_SUTUN, "Giriş")

    u.bosluk(ws, 4, 16.0)

    # --- Giris karti (B5:F11) --------------------------------------------
    u.kart(ws, 5, C_ETIKET, 11, C_ALAN_SON)
    u.bosluk(ws, 5, 14.0)
    u.sayfa_basligi(ws, 6, C_ETIKET, C_ALAN_SON, "Proje Öneri Formu",
                    "İyileştirme öneriniz birkaç dakikada gönderilir.")
    u.bosluk(ws, 8, 12.0)
    u.bosluk(ws, 9, 6.0)
    u.bosluk(ws, 10, 48.0)      # "Sisteme Gir" dugmesi buraya oturur
    u.bosluk(ws, 11, 14.0)

    # --- Kartin altindaki aciklama ---------------------------------------
    u.bosluk(ws, 12, 14.0)
    u.bilgi_kutusu(
        ws, 13, C_ETIKET, 13, C_ALAN_SON,
        "Gönderim sonrasında size bir öneri numarası verilir; önerinizin "
        "durumunu bu numara ve sicil numaranızla ProjeTakip dosyasından "
        "sorgulayabilirsiniz.",
    )
    ws.row_dimensions[13].height = 36.0

    u.bosluk(ws, 14, 12.0)
    u.not_satiri(ws, 15, C_ETIKET, C_ALAN_SON,
                 "Dosya açıldığında üstte bir güvenlik uyarısı çıkarsa "
                 "“İçeriği Etkinleştir” düğmesine basın.")

    ws.sheet_state = "visible"
    return ws


def _form_sayfasi(wb):
    """Oneri formu: uc numarali bolum, her bolum kendi nefes payinda.

    Form zeminin uzerinde durur, alanlar beyazdir: yazilabilecek yer ile
    okunacak yer birbirinden renkle ayrilir, cerceve cizmeye gerek kalmaz.
    """
    ws = wb.create_sheet(SAYFA_FORM)
    u.sayfa_hazirla(ws, SUTUNLAR, son_satir=34)
    u.masthead(ws, SON_SUTUN, "Öneri Formu")

    u.bosluk(ws, 4, 12.0)

    # Durum bandi: gonderim sonucunu formu terk etmeden gosterir.
    ws.row_dimensions[5].height = 26.0
    bant = u.birlestir(ws, 5, C_ETIKET, 5, C_ALAN_SON - 1)
    bant.alignment = u.hiza("left", "center", girinti=1)
    wb.defined_names.add(_ad("frm_bant", SAYFA_FORM, C_ETIKET, 5))

    u.bosluk(ws, 6, 12.0)

    # --- 1. Kimlik --------------------------------------------------------
    u.bolum_basligi(ws, 7, C_ETIKET, C_ALAN_SON, "1  ·  PERSONEL BİLGİSİ")
    u.bosluk(ws, 8, 6.0)
    u.bosluk(ws, 9, 10.0)

    _alan(ws, wb, 10, "Ad Soyad", "frm_ad_soyad", zorunlu=True)
    u.bosluk(ws, 11, 6.0)
    _alan(ws, wb, 12, "Sicil No", "frm_sicil_no", zorunlu=True)

    u.bosluk(ws, 13, 20.0)

    # --- 2. Mevcut durum --------------------------------------------------
    u.bolum_basligi(ws, 14, C_ETIKET, C_ALAN_SON, "2  ·  ŞU ANKİ DURUM",
                    "Neyin iyileştirilmesi gerektiğini anlatın.")
    u.bosluk(ws, 16, 10.0)

    _alan(ws, wb, 17, "Mevcut Durum", "frm_mevcut", zorunlu=True,
          yukseklik=72.0, coklu=True)
    u.ipucu_satiri(ws, 18, C_ALAN_BAS, C_ALAN_SON,
                   "Örnek: “Aynı dilekçe hem taranıyor hem e-postayla gönderiliyor.”")

    u.bosluk(ws, 19, 20.0)

    # --- 3. Oneri ---------------------------------------------------------
    u.bolum_basligi(ws, 20, C_ETIKET, C_ALAN_SON, "3  ·  ÖNERİNİZ",
                    "Nasıl daha iyi olabileceğini ve ne kazandıracağını yazın.")
    u.bosluk(ws, 22, 10.0)

    _alan(ws, wb, 23, "Öneri Başlığı", "frm_baslik", zorunlu=True)
    u.ipucu_satiri(ws, 24, C_ALAN_BAS, C_ALAN_SON,
                   "Örnek: “Kredi dosyası için tek sayfalık kontrol listesi”")
    _alan(ws, wb, 25, "Çözüm Öneriniz", "frm_cozum", zorunlu=True,
          yukseklik=72.0, coklu=True)
    u.ipucu_satiri(ws, 26, C_ALAN_BAS, C_ALAN_SON,
                   "Örnek: “Dilekçe yalnızca taransın, e-posta gönderilmesin.”")
    _alan(ws, wb, 27, "Beklenen Fayda", "frm_fayda", yukseklik=56.0, coklu=True)
    u.ipucu_satiri(ws, 28, C_ALAN_BAS, C_ALAN_SON,
                   "Örnek: “Dosya başına 5 dakika kazanç.”")

    # --- Dugmeler ---------------------------------------------------------
    u.bosluk(ws, 29, 16.0)
    u.ayrac(ws, 30, C_ETIKET, C_ALAN_SON)
    u.bosluk(ws, 31, 14.0)
    u.bosluk(ws, 32, 44.0)      # Gonder / Temizle dugmeleri
    u.bosluk(ws, 33, 14.0)

    ws.sheet_state = "veryHidden"
    return ws


def _ad(ad, sayfa, sutun, satir):
    from openpyxl.utils import get_column_letter
    from openpyxl.workbook.defined_name import DefinedName
    return DefinedName(ad, attr_text=f"'{sayfa}'!${get_column_letter(sutun)}${satir}")


def _alan(ws, wb, satir, etiket_metni, ad, zorunlu=False, liste=None,
          yukseklik=None, coklu=False):
    """Bir etiket + giris alani satiri kurar."""
    e = u.etiket(ws, satir, C_ETIKET, etiket_metni, zorunlu=zorunlu)
    if coklu:
        e.alignment = u.hiza("left", "top")

    u.form_alani(ws, wb, satir, C_ALAN_BAS, C_ALAN_SON, ad, SAYFA_FORM,
                 yukseklik=yukseklik, coklu_satir=coklu)

    if liste:
        dv = DataValidation(type="list", formula1=f"={liste}", allow_blank=True,
                            showDropDown=False, showErrorMessage=True,
                            errorTitle="Geçersiz seçim",
                            error="Lütfen listeden bir değer seçin.")
        ws.add_data_validation(dv)
        dv.add(u.adres(satir, C_ALAN_BAS, satir, C_ALAN_SON))


def kitap_uret(hedef_yol, listeler):
    wb = Workbook()
    wb.remove(wb.active)

    giris = _giris_sayfasi(wb)
    _form_sayfasi(wb)
    u.listeler_sayfasi_kur(wb, listeler)

    # Sekme rengi: kitap sekmeleri de ekranlarla ayni renk dilini konussun.
    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = RENK["DERIN_LACIVERT"]

    wb.active = wb.index(giris)
    wb.save(hedef_yol)
    return hedef_yol
