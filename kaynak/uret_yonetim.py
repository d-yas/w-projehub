# -*- coding: utf-8 -*-
"""ProjeYonetim.xlsm -- Değerlendirme ekibinin kitabinin sayfa kurulumu.

Dort ekran: Giris, Pano, Liste, Degerlendirme. Pano hem kitaptaki ILK
sayfadir hem de sifre girildikten sonra acilan ekrandir; Liste bir dugme
uzaktadir.

Ayrica dort gizli calisma sayfasi. Ikisi VERI DEPOSUDUR ve sistemin butun
verisini tasir -- bu kitap yalnizca bir ekran degil, ayni zamanda veritabani:

    Oneriler   satir basina bir gonderim (degismez)
    Olaylar    satir basina bir degerlendirme (yalnizca eklenir)

Digerleri turetilmis, her yenilemede bastan yazilan onbelleklerdir:

    Veri       konsolide tablo (liste ve pano bundan okur)
    PanoVeri   grafiklerin kaynak araliklari

Sutun duzeni modDepo.bas'taki O_* / E_* sabitleriyle ayni olmak zorundadir;
test_uretim.py basliklarin sirasini o sabitlerle karsilastirir.
"""

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

import uret_ortak as u
from tasarim import DURUM_RENK, FONT_BASLIK_AILE, KART, PT, RENK

SAYFA_GIRIS = "Giriş"
SAYFA_LISTE = "Liste"
SAYFA_DEGERLENDIRME = "Değerlendirme"
SAYFA_PANO = "Pano"
SAYFA_VERI = "Veri"
SAYFA_PANOVERI = "PanoVeri"
SAYFA_ONERILER = "Oneriler"
SAYFA_OLAYLAR = "Olaylar"

GIRIS_SUTUNLAR = [2.2, 24.0, 20.0, 20.0, 20.0, 14.0, 2.2]

# --- Liste duzeni (modKonsolide.bas'taki sabitlerle ayni olmali) ---------
# Sutun genislikleri en uzun degerlerine gore olculmustur: oneri numarasi
# ("PRJ-2026-0001") ve durum ("Standartlaştırıldı") kirpilmamalidir.
LISTE_SUTUNLAR = [2.2, 18.0, 13.0, 24.0, 54.0, 20.0, 2.2]
LISTE_BASLIK_SATIR = 8           # modKonsolide.FiltreKur: ILK_SATIR - 1
LISTE_ILK_SATIR = 9
LISTE_STIL_SATIR = 500           # bu kadar satir onceden bicimlendirilir
LISTE_ILK_SUTUN = 2              # B
LISTE_SON_SUTUN = 6              # F

# --- Degerlendirme duzeni ------------------------------------------------
DEG_SUTUNLAR = [2.2, 24.0, 24.0, 24.0, 20.0, 26.0, 2.2]
DEG_GECMIS_BASLIK = 33
DEG_GECMIS_ILK = 34              # modDegerlendirme.GECMIS_ILK_SATIR
DEG_GECMIS_ADET = 14

# --- Pano duzeni ---------------------------------------------------------
# Dort esit KPI karti tek sirada; aralarindaki ince sutunlar "oluk" gorevi
# gorur. Kartlar hucre birlestirmeyle degil bu izgarayla hizalanir.
PANO_SUTUNLAR = [2.2, 15.0, 15.0, 1.8, 15.0, 15.0, 1.8,
                 15.0, 15.0, 1.8, 15.0, 15.0, 2.2]
PANO_KART_SUTUNLARI = [(2, 3), (5, 6), (8, 9), (11, 12)]   # B/C E/F H/I K/L
PANO_KART_SATIR = 9                                        # 9-10-11
PANO_GRAFIK_SUTUNLARI = [(2, 6), (8, 12)]                  # B..F ve H..L
PANO_GRAFIK_BASLIK = 13
PANO_GRAFIK_UST = 15
PANO_GRAFIK_ALT = 31

# (etiket, adlandirilmis aralik, alt metin, vurgu rengi)
PANO_KARTLARI = [
    ("Toplam öneri", "pano_toplam", "sisteme gelen tüm öneriler",
     RENK["ANA_LACIVERT"]),
    ("Değerlendirme bekleyen", "pano_bekleyen", "Yeni + Değerlendirmede",
     RENK["VURGU_KEHRIBAR"]),
    ("Uygulamaya geçmiş", "pano_uygulanan", "pilot, ölçüm ve standart",
     RENK["ONAY_YAZI"]),
    ("Bu ay gelen", "pano_bu_ay", "içinde bulunduğumuz ay",
     RENK["KURUMSAL_MAVI"]),
]


def _ad(ad, sayfa, sutun, satir):
    return DefinedName(ad, attr_text=f"'{sayfa}'!${get_column_letter(sutun)}${satir}")


# ==========================================================================
#  Giris
# ==========================================================================
def _giris(wb):
    ws = wb.create_sheet(SAYFA_GIRIS)
    u.sayfa_hazirla(ws, GIRIS_SUTUNLAR, son_satir=30)
    u.masthead(ws, 7, "Giriş")

    u.bosluk(ws, 4, 16.0)

    # --- Giris karti (B5:F11) --------------------------------------------
    u.kart(ws, 5, 2, 11, 6)
    u.bosluk(ws, 5, 14.0)
    u.sayfa_basligi(ws, 6, 2, 6, "Proje Öneri Yönetimi",
                    "Öneriler burada değerlendirilir, önceliklendirilir ve izlenir.")
    u.bosluk(ws, 8, 12.0)
    u.bosluk(ws, 9, 6.0)
    u.bosluk(ws, 10, 48.0)      # "Sisteme Gir" dugmesi
    u.bosluk(ws, 11, 14.0)

    # Uyari bandi: kitap yazma kipinde acildiysa (kilit birakilamadi) burada
    # gorunur. Sorunun tek belirtisi budur -- personel o sirada oneri
    # gonderemez ama bunu ekip gormez.
    bant = u.birlestir(ws, 12, 2, 12, 6)
    bant.alignment = u.hiza("left", "center", girinti=1)
    ws.row_dimensions[12].height = 22.0
    wb.defined_names.add(_ad("giris_bant", SAYFA_GIRIS, 2, 12))

    u.bilgi_kutusu(
        ws, 13, 2, 13, 6,
        "Bu ekran yalnızca Değerlendirme ekibi içindir. Önerilerin ve "
        "değerlendirmelerin tamamı bu dosyanın içinde saklanır; her açılışta "
        "yanındaki yedek klasörüne günlük bir kopya alınır.",
    )
    ws.row_dimensions[13].height = 44.0

    u.bosluk(ws, 14, 12.0)
    u.not_satiri(ws, 15, 2, 6,
                 "Dosya açıldığında üstte bir güvenlik uyarısı çıkarsa "
                 "“İçeriği Etkinleştir” düğmesine basın.")

    ws.sheet_state = "visible"
    return ws


# ==========================================================================
#  Liste
# ==========================================================================
def _liste(wb):
    ws = wb.create_sheet(SAYFA_LISTE)
    # Satir/sutun basliklari BU ekranda acik kalir: burada gercek bir tablo
    # taranir, kullanici satir numarasina ve kaydirma cubuguna bakar.
    son_sutun = u.sayfa_hazirla(ws, LISTE_SUTUNLAR,
                                son_satir=LISTE_ILK_SATIR + LISTE_STIL_SATIR,
                                baslik_gizle=False)
    u.masthead(ws, son_sutun, "Liste")

    u.bosluk(ws, 4, 12.0)
    u.bosluk(ws, 5, 44.0)       # dugme seridi
    u.bosluk(ws, 6, 10.0)

    ozet = u.birlestir(ws, 7, LISTE_ILK_SUTUN, 7, LISTE_SON_SUTUN)
    ozet.font = u.yazi(RENK["METIN_GRI"], PT["OZET"])
    ozet.alignment = u.hiza("left", "center", girinti=1)
    ozet.value = "Önerileri görmek için “Önerileri Yenile” düğmesine basın."
    ws.row_dimensions[7].height = 18.0
    wb.defined_names.add(_ad("liste_ozet", SAYFA_LISTE, LISTE_ILK_SUTUN, 7))

    durum_sutun = LISTE_SON_SUTUN
    u.tablo_basligi(ws, LISTE_BASLIK_SATIR, LISTE_ILK_SUTUN, LISTE_SON_SUTUN,
                    ["Öneri No", "Tarih", "Gönderen", "Öneri Başlığı", "Durum"],
                    ortali=(durum_sutun,))

    son = LISTE_ILK_SATIR + LISTE_STIL_SATIR - 1
    u.tablo_govde_stili(ws, LISTE_ILK_SATIR, son, LISTE_ILK_SUTUN,
                        LISTE_SON_SUTUN, ortali=(durum_sutun,))

    # Durum sutununa rozet renkleri (kosullu bicimlendirme).
    u.durum_kosullu_bicim(ws, f"F{LISTE_ILK_SATIR}:F2000")

    # Baslik satiri ve ilk iki sutun sabit kalsin.
    ws.freeze_panes = f"C{LISTE_ILK_SATIR}"
    ws.sheet_state = "veryHidden"
    return ws


# ==========================================================================
#  Degerlendirme
# ==========================================================================
def _degerlendirme(wb):
    ws = wb.create_sheet(SAYFA_DEGERLENDIRME)
    son_sutun = u.sayfa_hazirla(ws, DEG_SUTUNLAR, son_satir=52)
    u.masthead(ws, son_sutun, "Değerlendirme")

    u.bosluk(ws, 4, 12.0)
    u.bosluk(ws, 5, 44.0)       # dugmeler
    u.bosluk(ws, 6, 10.0)

    ws.row_dimensions[7].height = 24.0
    bant = u.birlestir(ws, 7, 2, 7, 6)
    bant.alignment = u.hiza("left", "center", girinti=1)
    wb.defined_names.add(_ad("dg_bant", SAYFA_DEGERLENDIRME, 2, 7))

    u.bosluk(ws, 8, 12.0)

    # --- Secili oneri (salt okunur) --------------------------------------
    u.bolum_basligi(ws, 9, 2, 6, "SEÇİLİ ÖNERİ",
                    "Listede bir satır seçip “Seçiliyi Değerlendir” düğmesine basın.")
    u.bosluk(ws, 11, 8.0)

    _okuma_alani(ws, wb, 12, 2, "Öneri No", "dg_oneri_no", 3, 3)
    _okuma_alani(ws, wb, 12, 5, "Gönderim Tarihi", "dg_tarih", 6, 6)
    _okuma_alani(ws, wb, 13, 2, "Gönderen", "dg_gonderen", 3, 6)
    _okuma_alani(ws, wb, 14, 2, "Öneri Başlığı", "dg_baslik", 3, 6)
    _okuma_alani(ws, wb, 15, 2, "Mevcut Durum", "dg_mevcut", 3, 6,
                 yukseklik=58.0, coklu=True)
    _okuma_alani(ws, wb, 16, 2, "Çözüm Önerisi", "dg_cozum", 3, 6,
                 yukseklik=58.0, coklu=True)
    _okuma_alani(ws, wb, 17, 2, "Beklenen Fayda", "dg_fayda", 3, 6,
                 yukseklik=46.0, coklu=True)

    u.bosluk(ws, 18, 20.0)

    # --- Degerlendirme girisleri -----------------------------------------
    u.bolum_basligi(ws, 19, 2, 6, "DEĞERLENDİRME",
                    "Her kayıt geçmişe yeni bir satır olarak eklenir; "
                    "önceki değerlendirmeler silinmez.")
    u.bosluk(ws, 21, 8.0)

    _giris_alani(ws, wb, 22, 2, "Yeni Durum", "dg_yeni_durum", 3, 4,
                 liste="lst_durum", zorunlu=True)
    u.bosluk(ws, 23, 8.0)
    _giris_alani(ws, wb, 24, 2, "Karar Notu / Yorum", "dg_not", 3, 6,
                 yukseklik=76.0, coklu=True)

    u.ipucu_satiri(ws, 25, 3, 6,
                   "Her kayıt geçmişe eklenir; reddedilen öneriler için "
                   "gerekçe zorunludur.")
    u.bosluk(ws, 26, 18.0)
    u.ayrac(ws, 27, 2, 6)
    u.bosluk(ws, 28, 16.0)
    u.bosluk(ws, 29, 14.0)

    # --- Gecmis -----------------------------------------------------------
    u.bolum_basligi(ws, 30, 2, 6, "DEĞERLENDİRME GEÇMİŞİ",
                    "En yeni kayıt üstte. Tam geçmiş bu dosyanın içinde saklanır.")
    u.bosluk(ws, 32, 6.0)

    u.tablo_basligi(ws, DEG_GECMIS_BASLIK, 2, 6,
                    ["Tarih", "Değerlendiren", "Durum", "Karar Notu / Yorum"],
                    ortali=(4,))
    u.tablo_govde_stili(ws, DEG_GECMIS_ILK, DEG_GECMIS_ILK + DEG_GECMIS_ADET - 1,
                        2, 6, ortali=(4,))

    # Not sutunu iki sutun genisliginde: yorumlar tek satirda okunabilsin.
    ws.merge_cells(start_row=DEG_GECMIS_BASLIK, start_column=5,
                   end_row=DEG_GECMIS_BASLIK, end_column=6)
    for r in range(DEG_GECMIS_ILK, DEG_GECMIS_ILK + DEG_GECMIS_ADET):
        ws.merge_cells(start_row=r, start_column=5, end_row=r, end_column=6)
    u.durum_kosullu_bicim(
        ws, f"D{DEG_GECMIS_ILK}:D{DEG_GECMIS_ILK + DEG_GECMIS_ADET - 1}")

    ws.sheet_state = "veryHidden"
    return ws


def _okuma_alani(ws, wb, satir, etiket_sutun, etiket_metni, ad, c1, c2,
                 yukseklik=None, coklu=False):
    """Salt okunur bilgi alani.

    Giris alanlarindan bilerek FARKLI gorunur: beyaz kutusu YOKTUR, zeminde
    durur ve altinda ince bir cizgi vardir. "Buraya yazamazsin" bilgisi kilit
    uyarisiyla degil gorunumle verilir; beyaz yuzey yalnizca yazilabilir
    alanlara ayrilmistir.
    """
    e = u.etiket(ws, satir, etiket_sutun, etiket_metni)
    if coklu:
        e.alignment = u.hiza("left", "top")
    ws.row_dimensions[satir].height = yukseklik or 24.0

    h = u.birlestir(ws, satir, c1, satir, c2)
    h.font = u.yazi(RENK["METIN_KOYU"], PT["GOVDE"], kalin=not coklu)
    h.alignment = u.hiza("left", "top" if coklu else "center",
                         kaydir=coklu, girinti=1)
    for c in range(c1, c2 + 1):
        hh = ws.cell(row=satir, column=c)
        hh.border = u.Border(bottom=u.kenar(RENK["CIZGI_GRI"]))
    wb.defined_names.add(_ad(ad, SAYFA_DEGERLENDIRME, c1, satir))
    return h


def _giris_alani(ws, wb, satir, etiket_sutun, etiket_metni, ad, c1, c2,
                 liste=None, yukseklik=None, coklu=False, bicim=None,
                 zorunlu=False):
    """Kilitsiz giris alani (sayfa korumali oldugu icin yalnizca bunlar yazilabilir)."""
    e = u.etiket(ws, satir, etiket_sutun, etiket_metni, zorunlu=zorunlu)
    if coklu:
        e.alignment = u.hiza("left", "top")

    h = u.form_alani(ws, wb, satir, c1, c2, ad, SAYFA_DEGERLENDIRME,
                     yukseklik=yukseklik, coklu_satir=coklu)
    if bicim:
        for c in range(c1, c2 + 1):
            ws.cell(row=satir, column=c).number_format = bicim

    if liste:
        dv = DataValidation(type="list", formula1=f"={liste}", allow_blank=True,
                            showDropDown=False, showErrorMessage=True,
                            errorTitle="Geçersiz seçim",
                            error="Lütfen listeden bir değer seçin.")
        ws.add_data_validation(dv)
        dv.add(u.adres(satir, c1, satir, c2))
    return h


# ==========================================================================
#  Pano
# ==========================================================================
def _pano(wb):
    ws = wb.create_sheet(SAYFA_PANO)
    son_sutun = u.sayfa_hazirla(ws, PANO_SUTUNLAR, son_satir=33)
    u.masthead(ws, son_sutun, "Pano")

    u.bosluk(ws, 4, 12.0)
    u.bosluk(ws, 5, 44.0)       # dugmeler
    u.bosluk(ws, 6, 10.0)

    g = u.birlestir(ws, 7, 2, 7, son_sutun - 1)
    g.font = u.yazi(RENK["METIN_GRI"], PT["NOT"])
    g.alignment = u.hiza("left", "center", girinti=1)
    ws.row_dimensions[7].height = 16.0
    wb.defined_names.add(_ad("pano_guncelleme", SAYFA_PANO, 2, 7))

    u.bosluk(ws, 8, 10.0)

    # KPI hucreleri: gostergelerin TEK dogruluk kaynagi. Kartin gorunur govdesi
    # COM asamasinda bunun ustune cizilen bir sekildir (bkz. com_ek_islem);
    # sekil kurulamazsa bu katman tek basina okunur kalir.
    for (c1, c2), (etiket, ad, alt, vurgu) in zip(PANO_KART_SUTUNLARI, PANO_KARTLARI):
        u.kpi_karti(ws, wb, PANO_KART_SATIR, c1, c2, etiket, ad, SAYFA_PANO,
                    alt_metin=alt, vurgu_hex=vurgu, yuzey=False)

    u.bosluk(ws, 12, 20.0)

    # --- Grafik basliklari ve alani --------------------------------------
    for (c1, c2), metin in zip(PANO_GRAFIK_SUTUNLARI,
                               ["DURUM DAĞILIMI", "SON 12 AYIN GÖNDERİM TRENDİ"]):
        h = u.birlestir(ws, PANO_GRAFIK_BASLIK, c1, PANO_GRAFIK_BASLIK, c2)
        h.value = metin
        h.font = u.yazi(RENK["ANA_LACIVERT"], PT["BOLUM_BASLIK"], kalin=True,
                        aile=FONT_BASLIK_AILE)
        h.alignment = u.hiza("left", "bottom", girinti=1)
    ws.row_dimensions[PANO_GRAFIK_BASLIK].height = 22.0
    u.bosluk(ws, 14, 6.0)

    # Panel yuzeyi de sekil olarak cizilir; burada yalnizca satir yuksekligi
    # ayarlanir, boylece grafik dikdortgeni hesaplanabilir.
    for r in range(PANO_GRAFIK_UST, PANO_GRAFIK_ALT + 1):
        ws.row_dimensions[r].height = 18.0

    ws.sheet_state = "veryHidden"
    return ws


# ==========================================================================
#  Gizli calisma sayfalari
# ==========================================================================

# --- VERI DEPOSU ----------------------------------------------------------
# Bu iki liste sutun duzeninin PYTHON tarafidir; VBA tarafi modDepo.bas'taki
# O_* ve E_* sabitleridir. Ikisi ayrilirsa gonderim yanlis sutuna yazilir ve
# hicbir sey hata vermez -- test_uretim.py bu yuzden her basligin konumunu
# karsilik gelen sabitle karsilastirir.
ONERILER_BASLIKLARI = [
    "sema", "oneri_no", "tarih",
    "ad_soyad", "sicil_no",
    "mevcut_durum",
    "oneri_basligi", "cozum_onerisi", "beklenen_fayda",
    "gonderen_bilgisayar", "gonderen_kullanici",
]

OLAYLAR_BASLIKLARI = [
    "sema", "oneri_no", "olay_tarihi",
    "yeni_durum", "karar_notu",
    "degerlendiren_kullanici", "degerlendiren_bilgisayar",
]

# --- Turetilmis onbellek --------------------------------------------------
VERI_BASLIKLARI = [
    "oneri_no", "tarih", "ad_soyad", "sicil_no",
    "mevcut_durum", "oneri_basligi", "cozum_onerisi", "beklenen_fayda",
    "durum",
    "ilk_olay", "son_olay", "degerlendiren", "karar_notu", "olay_sayisi",
    "gonderen_kullanici", "kaynak_satir",
]


def _depo_sayfasi(wb, ad, basliklar):
    """Veri deposu sayfasi: baslik satiri, metin bicimi, cok gizli.

    Sutunlar METIN bicimlidir. Aksi halde Excel "10045" sicil numarasini
    sayiya, "2026-09-08T10:11:51" tarihini tarihe cevirir; basindaki sifirlar
    ve saniye bilgisi sessizce kaybolur. VBA tarafi yazarken bicimi ayrica
    zorlar, buradaki hazirlik elle bakildiginda da dogru gorunmesi icindir.
    """
    ws = wb.create_sheet(ad)
    for i, baslik in enumerate(basliklar, start=1):
        h = ws.cell(row=1, column=i, value=baslik)
        h.font = u.yazi(RENK["BEYAZ"], PT["NOT"], kalin=True)
        h.fill = u.dolgu(RENK["ANA_LACIVERT"])
        harf = get_column_letter(i)
        ws.column_dimensions[harf].width = 24
        for satir in range(2, 202):
            ws[f"{harf}{satir}"].number_format = "@"
    ws.freeze_panes = "A2"
    ws.sheet_state = "veryHidden"
    return ws


def _veri(wb):
    ws = wb.create_sheet(SAYFA_VERI)
    for i, ad in enumerate(VERI_BASLIKLARI, start=1):
        h = ws.cell(row=1, column=i, value=ad)
        h.font = u.yazi(RENK["BEYAZ"], PT["NOT"], kalin=True)
        h.fill = u.dolgu(RENK["ANA_LACIVERT"])
        ws.column_dimensions[get_column_letter(i)].width = 20
    sayac_sutun = len(VERI_BASLIKLARI) + 2
    ws.cell(row=1, column=sayac_sutun, value="adet")
    ws.cell(row=2, column=sayac_sutun, value=0)
    wb.defined_names.add(_ad("veri_adet", SAYFA_VERI, sayac_sutun, 2))
    ws.sheet_state = "veryHidden"
    return ws


def _panoveri(wb):
    ws = wb.create_sheet(SAYFA_PANOVERI)
    for sutun, baslik in ((1, "durum"), (2, "adet"), (4, "ay"), (5, "adet")):
        h = ws.cell(row=1, column=sutun, value=baslik)
        h.font = u.yazi(RENK["METIN_GRI"], PT["NOT"], kalin=True)
        ws.column_dimensions[get_column_letter(sutun)].width = 22

    # Etiket sutunlari METIN olarak biçimlendirilir. Aksi halde Excel "Mar 26"
    # gibi ay etiketlerini tarih sanip "26.Mar" diye yeniden yazar ve grafiğin
    # x ekseninde bazı aylar farklı görünür.
    for sutun in (1, 4):
        harf = get_column_letter(sutun)
        for satir in range(2, 16):
            ws[f"{harf}{satir}"].number_format = "@"

    ws.sheet_state = "veryHidden"
    return ws


# ==========================================================================
#  Dugmeler ve COM ek islemleri
# ==========================================================================
DUGMELER = [
    {"sayfa": SAYFA_GIRIS, "hucre": "C10", "metin": "Sisteme Gir  →",
     "makro": "SistemeGir", "varyant": "birincil"},

    {"sayfa": SAYFA_LISTE, "hucre": "B5", "metin": "⟳  Önerileri Yenile",
     "makro": "OnerileriYenile", "varyant": "birincil", "genislik": 160.0},
    {"sayfa": SAYFA_LISTE, "hucre": "B5", "metin": "Seçiliyi Değerlendir",
     "makro": "SeciliyiDegerlendir", "varyant": "ikincil", "genislik": 168.0,
     "sol_kaydir": 170.0},
    {"sayfa": SAYFA_LISTE, "hucre": "B5", "metin": "Pano",
     "makro": "PanoyaGit", "varyant": "ikincil", "genislik": 92.0,
     "sol_kaydir": 348.0},
    {"sayfa": SAYFA_LISTE, "hucre": "B5", "metin": "Çıkış",
     "makro": "Cikis", "varyant": "sessiz", "genislik": 80.0,
     "sol_kaydir": 450.0},

    {"sayfa": SAYFA_DEGERLENDIRME, "hucre": "B5",
     "metin": "✓  Değerlendirmeyi Kaydet", "makro": "DegerlendirmeKaydet",
     "varyant": "birincil", "genislik": 196.0},
    {"sayfa": SAYFA_DEGERLENDIRME, "hucre": "B5", "metin": "←  Listeye Dön",
     "makro": "ListeyeDon", "varyant": "ikincil", "genislik": 140.0,
     "sol_kaydir": 206.0},

    {"sayfa": SAYFA_PANO, "hucre": "B5", "metin": "⟳  Panoyu Yenile",
     "makro": "PanoyuYenile", "varyant": "birincil", "genislik": 156.0},
    {"sayfa": SAYFA_PANO, "hucre": "B5", "metin": "←  Listeye Dön",
     "makro": "ListeyeDon", "varyant": "ikincil", "genislik": 140.0,
     "sol_kaydir": 166.0},
    {"sayfa": SAYFA_PANO, "hucre": "B5", "metin": "Çıkış",
     "makro": "Cikis", "varyant": "sessiz", "genislik": 80.0,
     "sol_kaydir": 316.0},
]


def com_ek_islem(wb):
    """Pano'nun sekil katmanini ve grafiklerini kurar.

    openpyxl yuvarlak kose, yumusak golge ve canli grafik uretemez; bu is
    gercek Excel uzerinden yapilir.

    SIRA ONEMLIDIR: sekiller ve grafikler ayni cizim katmanindadir ve ekleme
    sirasina gore ust uste binerler. Panel sekli grafikten ONCE eklenmezse
    grafik panelin altinda kalir ve gorunmez.
    """
    import com_kurulum as com

    pano = wb.Worksheets(SAYFA_PANO)
    eski = pano.Visible
    pano.Visible = -1
    try:
        # 1) KPI kartlari -- hucre katmaninin ustune
        try:
            for (c1, c2), (etiket, ad, alt, vurgu) in zip(PANO_KART_SUTUNLARI,
                                                          PANO_KARTLARI):
                com.kpi_karti_ciz(pano, PANO_KART_SATIR, c1, c2,
                                  etiket, alt, vurgu, ad)
        except Exception:
            pass          # sekil katmani olmadan da gostergeler okunur

        # 2) Grafik panelleri -- grafiklerden ONCE
        try:
            for (c1, c2) in PANO_GRAFIK_SUTUNLARI:
                sol, ust, gen, yuk = com.hucre_kutusu(
                    pano, PANO_GRAFIK_UST, c1, PANO_GRAFIK_ALT, c2,
                    yatay_inset=KART["YATAY_INSET"])
                com.kart_govdesi(pano, sol, ust, gen, yuk, f"panel_{c1}")
        except Exception:
            pass

        # 3) Grafikler -- panelin uzerine
        _grafikleri_kur(wb, pano)
    finally:
        pano.Visible = eski


def _grafikleri_kur(wb, pano):
    """Iki grafik: durum dagilimi solda, 12 aylik trend sagda.

    Etiket ve deger araliklari AYRI AYRI verilir; seri elle kurulur.
    SetSourceData'nin otomatik tahmini guvenilir degil: bir aralikta ilk veri
    satirini baslik sanip kategoriyi dusuruyor, digerinde etiket sutununu
    ikinci bir seri sanip kategori adlarini 1, 2, 3'e ceviriyordu.
    """
    import com_kurulum as com
    from tasarim import rgb_long

    XL_BAR = 57            # xlBarClustered
    XL_LINE_MARKERS = 65   # xlLineMarkers
    XL_CATEGORY, XL_VALUE = 1, 2
    XL_TICK_NONE = -4142
    XL_LABEL_OUTSIDE_END = 2
    XL_MARKER_CIRCLE = 8

    # Durum grafiginin cubuklari listedeki rozetlerle AYNI renk ailesini
    # kullanir; boylece iki ekran ayni renk dilini konusur.
    durum_renkleri = [DURUM_RENK[d][1] for d in
                      ["Yeni", "Değerlendirmede", "Planlandı", "Pilot Uygulamada",
                       "Ölçümleniyor", "Standartlaştırıldı", "Beklemede",
                       "Reddedildi"]]

    tanimlar = [
        (PANO_GRAFIK_SUTUNLARI[0], "$A$2:$A$9", "$B$2:$B$9", XL_BAR,
         RENK["ANA_LACIVERT"], True, durum_renkleri),
        (PANO_GRAFIK_SUTUNLARI[1], "$D$2:$D$13", "$E$2:$E$13", XL_LINE_MARKERS,
         RENK["KURUMSAL_MAVI"], False, None),
    ]

    pv = wb.Worksheets(SAYFA_PANOVERI)
    ic = KART["YATAY_INSET"] + 10.0

    for (c1, c2), etiket_aralik, deger_aralik, tur, renk, ters, nokta_renkleri \
            in tanimlar:
        sol, ust, gen, yuk = com.hucre_kutusu(
            pano, PANO_GRAFIK_UST, c1, PANO_GRAFIK_ALT, c2,
            yatay_inset=ic, dikey_inset=10.0)

        co = pano.ChartObjects().Add(sol, ust, gen, yuk)
        co.Name = f"grafik_{c1}"
        ch = co.Chart
        ch.ChartType = tur

        try:
            co.ShapeRange.Line.Visible = 0
            co.ShapeRange.Fill.Visible = 0
        except Exception:
            pass

        while ch.SeriesCollection().Count > 0:
            ch.SeriesCollection(1).Delete()
        yeni = ch.SeriesCollection().NewSeries()
        yeni.Values = pv.Range(deger_aralik)
        yeni.XValues = pv.Range(etiket_aralik)

        try:
            # Grafik zemini SEFFAF: altindaki yuvarlak kose panel gorunsun.
            ch.ChartArea.Format.Fill.Visible = 0
            ch.ChartArea.Format.Line.Visible = 0
            ch.ChartArea.Font.Name = "Segoe UI"
            ch.ChartArea.Font.Size = PT["GRAFIK"]
            ch.ChartArea.Font.Color = rgb_long(RENK["METIN_GRI"])
            ch.PlotArea.Format.Fill.Visible = 0
        except Exception:
            pass

        try:
            seri = ch.SeriesCollection(1)
            seri.Format.Fill.ForeColor.RGB = rgb_long(renk)
            seri.Format.Line.ForeColor.RGB = rgb_long(renk)
        except Exception:
            pass

        if nokta_renkleri:                      # --- durum dagilimi ---------
            try:
                ch.ChartGroups(1).GapWidth = 55
                seri.Format.Line.Visible = 0
                seri.HasDataLabels = True
                etiketler = seri.DataLabels()
                etiketler.Position = XL_LABEL_OUTSIDE_END
                etiketler.Font.Size = PT["GRAFIK"]
                etiketler.Font.Name = "Segoe UI"
                etiketler.Font.Bold = True
                etiketler.Font.Color = rgb_long(RENK["METIN_KOYU"])
                for i, hex_kod in enumerate(nokta_renkleri, start=1):
                    seri.Points(i).Format.Fill.ForeColor.RGB = rgb_long(hex_kod)
            except Exception:
                pass
            try:
                # Deger ekseni ve izgarasi kaldirilir: sayiyi etiketler zaten
                # soyluyor, eksen + izgara yalnizca gurultu ekliyordu.
                # Izgara AYRICA kapatilmali; ekseni silmek onu goturmuyor.
                ch.Axes(XL_VALUE).HasMajorGridlines = False
                ch.Axes(XL_VALUE).HasMinorGridlines = False
                ch.Axes(XL_VALUE).Delete()
            except Exception:
                pass
        else:                                   # --- 12 aylik trend ---------
            try:
                seri.Format.Line.Weight = 2.25
                seri.Smooth = False
                seri.MarkerStyle = XL_MARKER_CIRCLE
                seri.MarkerSize = 5
                seri.MarkerBackgroundColor = rgb_long(RENK["BEYAZ"])
                seri.MarkerForegroundColor = rgb_long(renk)
                seri.HasDataLabels = False      # 12 etiket okunaksiz yapiyordu
            except Exception:
                pass
            try:
                deger = ch.Axes(XL_VALUE)
                deger.MinimumScale = 0
                deger.MajorUnit = 1
                deger.TickLabels.NumberFormat = "0"
                deger.TickLabels.Font.Size = PT["GRAFIK"]
                deger.Format.Line.Visible = 0
                deger.MajorTickMark = XL_TICK_NONE
                deger.MajorGridlines.Format.Line.ForeColor.RGB = \
                    rgb_long(RENK["CIZGI_INCE"])
                deger.MajorGridlines.Format.Line.Weight = 0.75
            except Exception:
                pass

        try:
            kategori = ch.Axes(XL_CATEGORY)
            kategori.TickLabels.Font.Size = PT["GRAFIK"]
            kategori.MajorTickMark = XL_TICK_NONE
            kategori.Format.Line.ForeColor.RGB = rgb_long(RENK["CIZGI_GRI"])
            if ters:
                kategori.ReversePlotOrder = True
        except Exception:
            pass

        # Baslik ve gosterge EN SONDA kapatilir: ChartType ve seri cagrilari
        # bunlari yeniden acabiliyor.
        try:
            ch.HasLegend = False
            ch.HasTitle = False
        except Exception:
            pass


# ==========================================================================
def kitap_uret(hedef_yol, listeler):
    wb = Workbook()
    wb.remove(wb.active)

    # Sayfa SIRASI ekrandaki sekme sirasidir: Pano en basta durur, oturum
    # acildiginda ilk gorunen sekme odur. Giris ekrani ikinci sirada olsa da
    # kitap acildiginda gorunur olan tek sayfa yine odur.
    _pano(wb)
    giris = _giris(wb)
    _liste(wb)
    _degerlendirme(wb)
    _depo_sayfasi(wb, SAYFA_ONERILER, ONERILER_BASLIKLARI)
    _depo_sayfasi(wb, SAYFA_OLAYLAR, OLAYLAR_BASLIKLARI)
    _veri(wb)
    _panoveri(wb)

    u.listeler_sayfasi_kur(wb, dict(listeler))

    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = RENK["DERIN_LACIVERT"]

    # Etkin sayfa gizli olamaz: kitap acilirken gorunur olan tek sayfa Giris.
    wb.active = wb.index(giris)

    wb.save(hedef_yol)
    return hedef_yol
